import datetime
import io
import openpyxl
import pytest
from django.contrib.auth import get_user_model
from django.test import Client
from django.urls import reverse
from django.utils import timezone

from work_time_reporter.exports import ExportService
from work_time_reporter.models import TimeLog

User = get_user_model()


@pytest.fixture
def timesheet_with_data(db, engineer_user, active_task, draft_timesheet):
    """Create a weekly timesheet with logged hours and comments."""
    monday = datetime.date.fromisocalendar(draft_timesheet.year, draft_timesheet.week_number, 1)

    for i in range(5):
        TimeLog.objects.create(
            user=engineer_user,
            task=active_task,
            timesheet=draft_timesheet,
            date=monday + datetime.timedelta(days=i),
            hours=8.0,
            comment=f"Development task progress day {i+1}",
        )
    return draft_timesheet


@pytest.mark.django_db
class TestExportService:
    def test_generate_weekly_pdf_success(self, timesheet_with_data):
        """Should produce valid PDF bytes starting with %PDF."""
        pdf_bytes = ExportService.generate_weekly_pdf(timesheet_with_data)

        assert isinstance(pdf_bytes, bytes)
        assert len(pdf_bytes) > 1000
        # Standard PDF magic header
        assert pdf_bytes.startswith(b"%PDF")

    def test_generate_weekly_pdf_empty_timesheet(self, draft_timesheet):
        """Should safely produce a PDF even when no hours are logged."""
        pdf_bytes = ExportService.generate_weekly_pdf(draft_timesheet)

        assert isinstance(pdf_bytes, bytes)
        assert pdf_bytes.startswith(b"%PDF")

    def test_generate_yearly_excel_success(self, engineer_user, active_task, draft_timesheet, timesheet_with_data):
        """Should produce a valid Excel workbook with three expected sheets."""
        excel_bytes = ExportService.generate_yearly_excel(engineer_user, draft_timesheet.year)

        assert isinstance(excel_bytes, bytes)
        assert len(excel_bytes) > 2000

        # Load into openpyxl to verify workbook structure and data
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        sheet_names = wb.sheetnames

        assert len(sheet_names) == 3
        # First sheet is Weekly Summary
        ws_summary = wb[sheet_names[0]]
        assert ws_summary.cell(row=1, column=1).value is not None

        # Third sheet contains detailed time logs
        ws_logs = wb[sheet_names[2]]
        assert ws_logs.max_row > 1


@pytest.mark.django_db
class TestWeeklyPDFExportView:
    def test_owner_can_export_pdf(self, engineer_client, draft_timesheet):
        """Owner of the timesheet should successfully download the PDF."""
        url = reverse('work_time_reporter:export_weekly_pdf', kwargs={'timesheet_id': draft_timesheet.id})
        response = engineer_client.get(url)

        assert response.status_code == 200
        assert response['Content-Type'] == 'application/pdf'
        assert 'attachment;' in response['Content-Disposition']
        assert f"W{draft_timesheet.week_number:02d}.pdf" in response['Content-Disposition']
        assert response.content.startswith(b"%PDF")

    def test_manager_can_export_pdf(self, manager_client, active_project, draft_timesheet):
        """Project Manager of subordinate should be permitted to download PDF."""
        url = reverse('work_time_reporter:export_weekly_pdf', kwargs={'timesheet_id': draft_timesheet.id})
        response = manager_client.get(url)

        assert response.status_code == 200
        assert response['Content-Type'] == 'application/pdf'
        assert response.content.startswith(b"%PDF")

    def test_unrelated_user_forbidden(self, other_user, draft_timesheet):
        """An unrelated user cannot export another user's timesheet."""
        client = Client()
        client.force_login(other_user)

        url = reverse('work_time_reporter:export_weekly_pdf', kwargs={'timesheet_id': draft_timesheet.id})
        response = client.get(url)

        assert response.status_code == 403

    def test_unauthenticated_redirects_to_login(self, draft_timesheet):
        """Anonymous user should be redirected to login page."""
        client = Client()
        url = reverse('work_time_reporter:export_weekly_pdf', kwargs={'timesheet_id': draft_timesheet.id})
        response = client.get(url)

        assert response.status_code == 302
        assert 'login' in response.url

    def test_not_found_returns_404(self, engineer_client):
        """Non-existent timesheet ID should return 404."""
        url = reverse('work_time_reporter:export_weekly_pdf', kwargs={'timesheet_id': 999999})
        response = engineer_client.get(url)

        assert response.status_code == 404


@pytest.mark.django_db
class TestYearlyExcelExportView:
    def test_user_can_export_own_yearly_excel(self, engineer_client):
        """User can export their own yearly report."""
        current_year = timezone.now().year
        url = reverse('work_time_reporter:export_yearly_excel', kwargs={'year': current_year})
        response = engineer_client.get(url)

        assert response.status_code == 200
        assert response['Content-Type'] == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        assert 'attachment;' in response['Content-Disposition']
        assert f"{current_year}.xlsx" in response['Content-Disposition']

    def test_export_yearly_excel_default_year(self, engineer_client):
        """Yearly export route without year param defaults to current year."""
        url = reverse('work_time_reporter:export_yearly_excel_current')
        response = engineer_client.get(url)

        assert response.status_code == 200
        assert response['Content-Type'] == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'

    def test_manager_can_export_subordinate_excel(self, manager_client, active_project, engineer_user):
        """Manager can export yearly report of team member via ?user_id= query param."""
        current_year = timezone.now().year
        url = reverse('work_time_reporter:export_yearly_excel', kwargs={'year': current_year})
        response = manager_client.get(f"{url}?user_id={engineer_user.id}")

        assert response.status_code == 200
        assert response['Content-Type'] == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        assert engineer_user.username in response['Content-Disposition']

    def test_unrelated_user_forbidden_subordinate_excel(self, other_user, engineer_user):
        """An unrelated user cannot export another user's annual report."""
        client = Client()
        client.force_login(other_user)

        current_year = timezone.now().year
        url = reverse('work_time_reporter:export_yearly_excel', kwargs={'year': current_year})
        response = client.get(f"{url}?user_id={engineer_user.id}")

        assert response.status_code == 403
